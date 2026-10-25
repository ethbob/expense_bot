"""Форматирование текстов, расчёт периодов и экспорт в Excel."""
import html
import io
import math
from datetime import datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .db import Database

MONTHS = ["январь", "февраль", "март", "апрель", "май", "июнь",
          "июль", "август", "сентябрь", "октябрь", "ноябрь", "декабрь"]

STATUS_LABEL = {
    "pending": "⏳ Не разобрана",
    "expense": "✅ Расход",
    "not_expense": "➖ Не расход",
}
STATUS_ICON = {"pending": "⏳", "expense": "✅", "not_expense": "➖"}


# ---------- общие помощники ----------

def fmt(amount: Decimal | int | float) -> str:
    value = Decimal(amount).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    return f"{value:,.2f}".replace(",", " ")


def short_addr(addr: str | None) -> str:
    if not addr:
        return ""
    return f"{addr[:6]}…{addr[-4:]}"


def month_title(year: int, month: int) -> str:
    return f"{MONTHS[month - 1].capitalize()} {year}"


def shift_month(year: int, month: int, delta: int) -> tuple[int, int]:
    total = year * 12 + (month - 1) + delta
    return total // 12, total % 12 + 1


def month_bounds(year: int, month: int, tz: ZoneInfo) -> tuple[datetime, datetime]:
    start = datetime(year, month, 1, tzinfo=tz)
    ny, nm = shift_month(year, month, 1)
    return start, datetime(ny, nm, 1, tzinfo=tz)


def local(dt: datetime, tz: ZoneInfo) -> datetime:
    return dt.astimezone(tz)


# ---------- карточка транзакции ----------

def tx_card(tx, tz: ZoneInfo) -> str:
    if tx["source"] == "chain":
        lines = ["💸 <b>Перевод USDT (BEP-20)</b>"]
    else:
        lines = ["✍️ <b>Ручной расход</b>"]
    lines.append(f"Сумма: <b>{fmt(tx['amount'])} USDT</b>")
    if tx["to_address"]:
        lines.append(f"Куда: <code>{tx['to_address']}</code>")
    lines.append(f"Время: {local(tx['occurred_at'], tz):%d.%m.%Y %H:%M}")
    lines.append(f"Статус: {STATUS_LABEL[tx['status']]}")
    if tx["note"]:
        lines.append(f"📝 {html.escape(tx['note'])}")
    if tx["tx_hash"]:
        lines.append(f'<a href="https://bscscan.com/tx/{tx["tx_hash"]}">Открыть в BscScan</a>')
    if tx["status"] == "pending":
        lines.append("\n<b>Это расход?</b>")
    return "\n".join(lines)


def tx_button_text(tx, tz: ZoneInfo) -> str:
    icon = STATUS_ICON[tx["status"]]
    if tx["source"] == "manual":
        icon += "✍️"
    label = tx["note"] or short_addr(tx["to_address"]) or "без комментария"
    text = f"{icon} {fmt(tx['amount'])} · {local(tx['occurred_at'], tz):%d.%m} · {label}"
    return text if len(text) <= 60 else text[:59] + "…"


# ---------- сводки ----------

def _summary_lines(s) -> list[str]:
    lines = [f"💰 Расходы: <b>{fmt(s['exp_total'])} USDT</b>"]
    if s["exp_count"]:
        lines.append(f"   • переводы: {fmt(s['chain_total'])} ({s['chain_count']} шт.)")
        lines.append(f"   • вручную: {fmt(s['manual_total'])} ({s['manual_count']} шт.)")
    if s["not_count"]:
        lines.append(f"➖ Не расход: {fmt(s['not_total'])} USDT ({s['not_count']} шт.)")
    if s["pend_count"]:
        lines.append(
            f"⏳ Не разобрано: {fmt(s['pend_total'])} USDT ({s['pend_count']} шт.) — в итог не входят"
        )
    return lines


async def tracking_since(db: Database, tz: ZoneInfo) -> datetime | None:
    raw = await db.get_kv("started_at")
    if not raw:
        return None
    return datetime.fromisoformat(raw).astimezone(tz)


async def month_summary(db: Database, tz: ZoneInfo, year: int, month: int) -> str:
    start, end = month_bounds(year, month, tz)
    now = datetime.now(tz)
    s = await db.summary(start, end)
    is_current = (year, month) == (now.year, now.month)

    title = f"📊 <b>{month_title(year, month)}</b>"
    if is_current:
        title += f" (01.{month:02d} — {now:%d.%m})"

    # средние считаем только по дням, когда учёт уже велся
    since = await tracking_since(db, tz)
    eff_start = max(start, since) if since else start
    eff_end = min(end, now)
    days = max(1, math.ceil((eff_end - eff_start).total_seconds() / 86400))

    lines = [title, ""] + _summary_lines(s)
    if s["exp_count"]:
        lines.append(f"📅 В среднем в день: {fmt(s['exp_total'] / days)} USDT")
    return "\n".join(lines)


async def period_summary(db: Database, tz: ZoneInfo, days: int) -> str:
    now = datetime.now(tz)
    start = now - timedelta(days=days)
    s = await db.summary(start, now + timedelta(minutes=1))

    lines = [f"📈 <b>За последние {days} дней</b>",
             f"{start:%d.%m.%Y} — {now:%d.%m.%Y}", ""]
    lines += _summary_lines(s)

    since = await tracking_since(db, tz)
    effective_days = days
    if since and since > start:
        effective_days = max(1, math.ceil((now - since).total_seconds() / 86400))
    if s["exp_count"]:
        per_day = s["exp_total"] / effective_days
        lines.append(f"📅 В среднем в день: {fmt(per_day)} USDT")
        if days > 30 and effective_days >= 30:
            lines.append(f"🗓 В среднем за 30 дней: {fmt(per_day * 30)} USDT")
    if since and since > start:
        lines.append(f"\nℹ️ Учёт ведётся с {since:%d.%m.%Y}, средние посчитаны с этой даты.")
    return "\n".join(lines)


# ---------- экспорт ----------

STATUS_XLSX = {"pending": "Не разобрана", "expense": "Расход", "not_expense": "Не расход"}


def build_xlsx(rows, tz: ZoneInfo) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Транзакции"
    headers = ["Дата", "Тип", "Сумма, USDT", "Статус", "Примечание", "Получатель", "Хеш транзакции"]
    ws.append(headers)
    header_fill = PatternFill("solid", fgColor="DDEBF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    total = Decimal(0)
    for r in rows:
        ws.append([
            local(r["occurred_at"], tz).replace(tzinfo=None),
            "Перевод" if r["source"] == "chain" else "Вручную",
            float(r["amount"]),
            STATUS_XLSX[r["status"]],
            r["note"] or "",
            r["to_address"] or "",
            r["tx_hash"] or "",
        ])
        if r["status"] == "expense":
            total += r["amount"]

    last = ws.max_row
    for row in ws.iter_rows(min_row=2, max_row=last):
        row[0].number_format = "DD.MM.YYYY HH:MM"
        row[2].number_format = "#,##0.00"

    ws.append([])
    ws.append(["Итого расходов", "", float(total)])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.cell(ws.max_row, 3).font = Font(bold=True)
    ws.cell(ws.max_row, 3).number_format = "#,##0.00"

    for col, width in zip("ABCDEFG", [17, 10, 14, 14, 36, 44, 68]):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
